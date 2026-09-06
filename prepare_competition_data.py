#!/usr/bin/env python3
"""把用户自行导出的评论快照转换为 LocalPulse 的比赛输入格式。

比赛模式只读取本地 CSV / JSON / XLSX，不登录平台、不访问网络，也不保留昵称、性别、IP、UID
或回复对象等非必要字段。实时网页接口脚本仍保留在 ``bili_comment.py``，但默认被禁用，
仅供已经取得明确授权的研究场景使用。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


OUTPUT_FIELDS = ("id", "platform", "content", "likes", "timestamp")

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "content": ("content", "comment", "评论", "评论内容", "text", "comment_text", "内容", "message"),
    "likes": ("likes", "like_count", "点赞数", "点赞数量", "点赞", "赞", "like"),
    "timestamp": ("timestamp", "time", "发布时间", "时间", "reply_time", "回复时间", "评论时间", "ctime"),
    "platform": ("platform", "source", "平台", "来源"),
}

PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
URL_RE = re.compile(r"(?i)\b(?:https?://|www\.)[^\s]+")


class ExportError(ValueError):
    """用户可理解的本地导出错误。"""


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    return re.sub(r"[\s_\-—:：()（）]+", "", text)


def _read_json(path: Path) -> list[dict[str, Any]]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise ExportError("JSON 必须使用 UTF-8 编码") from exc
    except json.JSONDecodeError as exc:
        raise ExportError(f"JSON 格式错误：第 {exc.lineno} 行") from exc

    if isinstance(payload, dict):
        for key in ("comments", "data", "records", "items"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
        else:
            payload = [payload]
    if not isinstance(payload, list):
        raise ExportError("JSON 顶层必须是对象或数组")
    return [item if isinstance(item, dict) else {"content": item} for item in payload]


def _read_csv(path: Path) -> list[dict[str, Any]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return list(csv.DictReader(handle))
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ExportError(f"CSV 编码无法识别：{last_error}")


def _read_xlsx(path: Path) -> tuple[list[dict[str, Any]], int]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 的异常信息对用户通常不够直观
        raise ExportError(f"XLSX 无法读取：{exc}") from exc

    rows: list[dict[str, Any]] = []
    sheet_count = 0
    try:
        for sheet in workbook.worksheets:
            values = sheet.iter_rows(values_only=True)
            try:
                header_row = next(values)
            except StopIteration:
                continue
            headers = [str(value).strip() if value is not None else "" for value in header_row]
            if not any(headers):
                continue
            sheet_count += 1
            for row in values:
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                rows.append({headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]})
    finally:
        workbook.close()
    return rows, sheet_count


def read_export(path: str | Path) -> tuple[list[dict[str, Any]], int | None]:
    source = Path(path).expanduser()
    if not source.exists():
        raise ExportError(f"输入文件不存在：{source}")
    if source.is_dir():
        raise ExportError(f"输入路径是目录，不是 CSV/JSON/XLSX 文件：{source}")
    suffix = source.suffix.lower()
    if suffix == ".json":
        return _read_json(source), None
    if suffix == ".csv":
        return _read_csv(source), None
    if suffix == ".xlsx":
        return _read_xlsx(source)
    raise ExportError("仅支持 CSV、JSON 或 XLSX 文件")


def _find_value(row: dict[str, Any], target: str) -> Any:
    normalized = {normalize_header(key): value for key, value in row.items()}
    for alias in FIELD_ALIASES[target]:
        key = normalize_header(alias)
        if key in normalized:
            return normalized[key]
    return None


def _parse_number(value: Any) -> int | float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return max(0, value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    multiplier = 1.0
    if text.endswith(("万", "万+")):
        text = text.rstrip("+").removesuffix("万")
        multiplier = 10000.0
    elif text.lower().endswith("k"):
        text = text[:-1]
        multiplier = 1000.0
    try:
        result = max(0.0, float(text) * multiplier)
    except ValueError:
        return None
    return int(result) if result.is_integer() else result


def _parse_timestamp(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime(value.year, value.month, value.day)
    elif isinstance(value, (int, float)) and value > 10**9:
        parsed = datetime.fromtimestamp(value, tz=timezone.utc).replace(tzinfo=None)
    else:
        text = str(value).strip().replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
        if not text:
            return None
        parsed = None
        for candidate in (text, text.replace(" ", "T")):
            try:
                parsed = datetime.fromisoformat(candidate)
                break
            except ValueError:
                continue
        if parsed is None:
            return None
    return parsed.isoformat(timespec="seconds")


def _sanitize_content(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub(r"\s+", " ", text).strip()
    text = PHONE_RE.sub("[手机号已脱敏]", text)
    text = EMAIL_RE.sub("[邮箱已脱敏]", text)
    text = URL_RE.sub("[链接已脱敏]", text)
    return text


def transform_rows(rows: Iterable[dict[str, Any]], platform: str) -> tuple[list[dict[str, Any]], dict[str, int]]:
    output: list[dict[str, Any]] = []
    seen: set[tuple[Any, ...]] = set()
    duplicates_removed = 0
    invalid_removed = 0

    for row in rows:
        content = _sanitize_content(_find_value(row, "content"))
        if not content:
            invalid_removed += 1
            continue
        likes = _parse_number(_find_value(row, "likes"))
        timestamp = _parse_timestamp(_find_value(row, "timestamp"))
        source_platform = str(_find_value(row, "platform") or platform).strip() or platform
        key = (source_platform, content, likes, timestamp)
        if key in seen:
            duplicates_removed += 1
            continue
        seen.add(key)
        output.append(
            {
                "id": f"bili-{len(output) + 1:06d}",
                "platform": source_platform,
                "content": content,
                "likes": likes,
                "timestamp": timestamp,
            }
        )

    return output, {"duplicates_removed": duplicates_removed, "invalid_removed": invalid_removed}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_export(records: list[dict[str, Any]], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(OUTPUT_FIELDS), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


def write_provenance(
    source: Path,
    output: Path,
    records: list[dict[str, Any]],
    stats: dict[str, int],
    input_rows: int,
    sheet_count: int | None,
    platform: str,
) -> Path:
    provenance_path = output.with_suffix(output.suffix + ".provenance.json")
    payload = {
        "schema_version": "1.0",
        "collection_mode": "user_provided_local_export",
        "live_crawler_enabled": False,
        "platform": platform,
        "source_file_name": source.name,
        "source_sha256": _sha256(source),
        "source_type": source.suffix.lower().lstrip("."),
        "source_sheet_count": sheet_count,
        "input_rows": input_rows,
        "valid_rows": len(records),
        **stats,
        "fields_kept": list(OUTPUT_FIELDS),
        "fields_dropped": ["nickname", "sex", "ip", "uid", "user_id", "reply_target", "user_level"],
        "content_redactions": ["phone", "email", "url"],
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "output_file_name": output.name,
        "output_sha256": _sha256(output),
        "compliance_note": "仅在已获得数据来源授权、且比赛主办方允许的范围内使用；本文件不证明上游数据本身已获授权。",
    }
    provenance_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return provenance_path


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="将本地 B 站评论导出快照转换为 LocalPulse 输入 CSV")
    parser.add_argument("--input", required=True, help="本地 CSV、JSON 或 XLSX 导出文件")
    parser.add_argument("--output", required=True, help="输出 LocalPulse CSV 路径")
    parser.add_argument("--platform", default="Bilibili", help="缺少平台列时使用的平台名，默认 Bilibili")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    source = Path(args.input).expanduser()
    output = Path(args.output).expanduser()
    try:
        rows, sheet_count = read_export(source)
        records, stats = transform_rows(rows, args.platform)
        if not records:
            raise ExportError("没有可导出的有效评论：请检查 content/评论内容 列")
        write_export(records, output)
        provenance = write_provenance(source, output, records, stats, len(rows), sheet_count, args.platform)
    except ExportError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    print(f"输入行数: {len(rows)}")
    print(f"有效行数: {len(records)}")
    print(f"去重行数: {stats['duplicates_removed']}")
    print(f"无效行数: {stats['invalid_removed']}")
    print(f"输出: {output}")
    print(f"来源记录: {provenance}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
